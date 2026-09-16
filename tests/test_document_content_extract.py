"""Tests for structural document extraction (app.documents.content_extract)."""

from __future__ import annotations

import pytest
from docx import Document
from pypdf import PdfWriter
from reportlab.pdfgen import canvas as rl_canvas

from app.documents.content_extract import (
    MAX_SOURCE_SIZE_BYTES,
    clear_extraction_cache,
    extract_document,
)
from app.errors import DocumentError


def test_txt_extraction(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("First paragraph.\n\nSecond paragraph.", encoding="utf-8")

    content = extract_document(path)

    assert content.doc_type == "txt"
    assert [b.kind for b in content.blocks] == ["paragraph", "paragraph"]
    assert content.blocks[0].text == "First paragraph."
    assert content.word_count == 4


def test_markdown_extraction_preserves_structure(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text(
        "# Title\n\nIntro paragraph.\n\n## Section\n\n- one\n- two\n\n1. first\n2. second\n",
        encoding="utf-8",
    )

    content = extract_document(path)

    kinds = [b.kind for b in content.blocks]
    assert kinds == ["heading", "paragraph", "heading", "list", "list"]
    assert content.title == "Title"
    assert content.blocks[3].ordered is False
    assert content.blocks[3].items == ["one", "two"]
    assert content.blocks[4].ordered is True
    assert content.blocks[4].items == ["first", "second"]


def test_docx_extraction_headings_lists_and_tables(tmp_path):
    doc = Document()
    doc.add_heading("Report", level=1)
    doc.add_paragraph("Overview text.")
    doc.add_paragraph("Bullet one", style="List Bullet")
    doc.add_paragraph("Bullet two", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "2"
    path = tmp_path / "report.docx"
    doc.save(path)

    content = extract_document(path)

    kinds = [b.kind for b in content.blocks]
    assert kinds == ["heading", "paragraph", "list", "table"]
    assert content.blocks[2].items == ["Bullet one", "Bullet two"]
    assert content.blocks[3].rows == [["A", "B"], ["1", "2"]]


def test_pdf_with_no_extractable_text(tmp_path):
    path = tmp_path / "doc.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(DocumentError, match="no extractable text"):
        extract_document(path)


def test_pdf_page_range_extracts_only_selected_pages(tmp_path):
    path = tmp_path / "three-pages.pdf"
    canvas = rl_canvas.Canvas(str(path))
    for page_number in range(1, 4):
        canvas.drawString(72, 700, f"Unique source page {page_number}")
        canvas.showPage()
    canvas.save()

    content = extract_document(path, page_range=(2, 2), use_cache=False)

    text = " ".join(block.text or "" for block in content.blocks)
    assert "source page 2" in text
    assert "source page 1" not in text
    assert "source page 3" not in text
    assert content.warnings == ["Included source pages 2–2 of 3."]


def test_pdf_page_range_is_validated(tmp_path):
    path = tmp_path / "one-page.pdf"
    canvas = rl_canvas.Canvas(str(path))
    canvas.drawString(72, 700, "One page")
    canvas.save()

    with pytest.raises(DocumentError, match="Invalid PDF page range"):
        extract_document(path, page_range=(2, 4), use_cache=False)


def test_extraction_cache_reuses_unchanged_content(tmp_path, monkeypatch):
    import app.documents.content_extract as extraction

    path = tmp_path / "cached.pdf"
    path.write_bytes(b"%PDF-test")
    calls = []

    def fake_extract(_path, page_range=None, progress=None):
        calls.append((page_range, progress))
        return [extraction.ContentBlock(kind="paragraph", text="Cached text")], "Title", []

    clear_extraction_cache()
    monkeypatch.setattr(extraction, "_extract_pdfium", fake_extract)

    first = extract_document(path)
    second = extract_document(path)

    assert first == second
    assert len(calls) == 1


def test_cache_has_a_byte_budget_and_isolates_callers(tmp_path, monkeypatch):
    import app.documents.content_extract as extraction

    clear_extraction_cache()
    monkeypatch.setattr(extraction, "MAX_EXTRACTION_CACHE_BYTES", 1500)
    paths = [tmp_path / f"source-{i}.txt" for i in range(3)]
    for path in paths:
        path.write_text("Synthetic report " * 25, encoding="utf-8")
        content = extract_document(path)
        content.blocks[0].text = "Mutated caller copy"
    assert sum(map(len, extraction._EXTRACTION_CACHE.values())) <= 1500
    assert len(extraction._EXTRACTION_CACHE) < 3
    assert extract_document(paths[-1]).blocks[0].text.startswith("Synthetic report")

    paths[-1].write_text("Changed source", encoding="utf-8")
    assert extract_document(paths[-1]).blocks[0].text == "Changed source"


def test_cache_does_not_keep_oversized_entries(tmp_path, monkeypatch):
    import app.documents.content_extract as extraction

    clear_extraction_cache()
    monkeypatch.setattr(extraction, "MAX_EXTRACTION_CACHE_BYTES", 100)
    path = tmp_path / "large.txt"
    path.write_text("Synthetic text " * 200, encoding="utf-8")
    assert extract_document(path).word_count == 400
    assert not extraction._EXTRACTION_CACHE


@pytest.mark.parametrize("backend", ["pdfium", "pypdf"])
def test_pdf_backends_preserve_numbers_and_handle_ranges(tmp_path, backend):
    path = tmp_path / "synthetic.pdf"
    canvas = rl_canvas.Canvas(str(path))
    canvas.setTitle("Synthetic report")
    for number in range(1, 4):
        canvas.drawString(72, 700, f"Section {number}: Revenue 1,234.56; return -12.5%")
        canvas.showPage()
    canvas.save()
    content = extract_document(path, pdf_backend=backend, page_range=(2, 25), use_cache=False)
    text = " ".join(b.text or "" for b in content.blocks)
    assert "Section 1" not in text
    assert "Section 2" in text and "Section 3" in text
    assert "1,234.56" in text and "-12.5%" in text
    assert content.title == "Synthetic report"
    assert content.warnings == ["Included source pages 2–3 of 3."]


@pytest.mark.parametrize("password", ["", "secret"])
def test_pdfium_rejects_encrypted_files(tmp_path, password):
    path = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt(password)
    with path.open("wb") as handle:
        writer.write(handle)
    with pytest.raises(DocumentError, match="password-protected"):
        extract_document(path)


def test_pdfium_rejects_corrupt_pdf(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a pdf")
    with pytest.raises(DocumentError, match="Failed to read PDF"):
        extract_document(path)


def test_missing_file(tmp_path):
    with pytest.raises(DocumentError, match="not found"):
        extract_document(tmp_path / "missing.txt")


def test_unsupported_extension(tmp_path):
    path = tmp_path / "sheet.xlsx"
    path.write_bytes(b"not a real spreadsheet")

    with pytest.raises(DocumentError, match=r"Unsupported document type: \.xlsx"):
        extract_document(path)


def test_oversized_file_rejected(tmp_path):
    path = tmp_path / "big.txt"
    path.write_bytes(b"a" * (MAX_SOURCE_SIZE_BYTES + 1))

    with pytest.raises(DocumentError, match="too large"):
        extract_document(path)


def test_empty_file_rejected(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_bytes(b"")

    with pytest.raises(DocumentError, match="empty"):
        extract_document(path)


def test_binary_txt_rejected(tmp_path):
    path = tmp_path / "binary.txt"
    path.write_bytes(b"\x00\x01\x02binarydata")

    with pytest.raises(DocumentError, match="does not look like text"):
        extract_document(path)


def test_unusual_filename_still_extracts(tmp_path):
    path = tmp_path / "client notes (draft) v2 — final.txt"
    path.write_text("Some content.", encoding="utf-8")

    content = extract_document(path)

    assert content.blocks[0].text == "Some content."


def test_extraction_never_modifies_source(tmp_path):
    path = tmp_path / "notes.txt"
    original = "Untouched content.\n\nSecond paragraph."
    path.write_text(original, encoding="utf-8")
    before_mtime = path.stat().st_mtime

    extract_document(path)

    assert path.read_text(encoding="utf-8") == original
    assert path.stat().st_mtime == before_mtime


def test_directory_rejected(tmp_path):
    directory = tmp_path / "a_directory.txt"
    directory.mkdir()

    with pytest.raises(DocumentError, match="Not a regular file"):
        extract_document(directory)
