"""Tests for branded PDF generation (app.pdf.service)."""

from __future__ import annotations

import pytest
from pypdf import PdfReader

from app.errors import DocumentError, PdfGenerationError
from app.pdf.service import create_branded_pdf


def test_creates_valid_verified_pdf(tmp_path):
    source = tmp_path / "notes.txt"
    source.write_text("Hello world.\n\nSecond paragraph.", encoding="utf-8")
    outputs = tmp_path / "outputs"

    result = create_branded_pdf(source, outputs)

    output_path = outputs / "notes-almond.pdf"
    assert result.output_path == str(output_path)
    assert output_path.is_file()
    assert output_path.stat().st_size > 0
    assert result.pages >= 1
    assert result.success is True

    reader = PdfReader(output_path)
    assert len(reader.pages) == result.pages
    assert len(reader.pages[0].images) >= 1  # Almond logo embedded


def test_pdf_source_with_real_text_is_converted(tmp_path):
    from reportlab.pdfgen import canvas as rl_canvas

    source = tmp_path / "source.pdf"
    canvas = rl_canvas.Canvas(str(source))
    canvas.drawString(72, 700, "Hello from a real PDF source document.")
    canvas.save()
    outputs = tmp_path / "outputs"

    result = create_branded_pdf(source, outputs)

    assert result.success is True
    assert result.pages >= 1
    text = PdfReader(result.output_path).pages[0].extract_text()
    assert "Hello from a real PDF source document." in text


def test_multi_page_document_produces_multiple_pages(tmp_path):
    source = tmp_path / "long.md"
    filler = "Filler paragraph text to force wrapping across pages. " * 40
    body = "\n\n".join(f"## Section {i}\n\n{filler}" for i in range(20))
    source.write_text(f"# Long Document\n\n{body}", encoding="utf-8")
    outputs = tmp_path / "outputs"

    result = create_branded_pdf(source, outputs)

    assert result.pages > 1


def test_headings_paragraphs_and_lists_are_rendered(tmp_path):
    source = tmp_path / "structured.md"
    source.write_text(
        "# Title\n\nIntro paragraph.\n\n## Section\n\n- item one\n- item two\n",
        encoding="utf-8",
    )
    outputs = tmp_path / "outputs"

    result = create_branded_pdf(source, outputs)
    text = PdfReader(result.output_path).pages[0].extract_text()

    assert "Title" in text
    assert "Intro paragraph." in text
    assert "Section" in text
    assert "item one" in text
    assert "item two" in text


def test_footer_and_page_numbers_present(tmp_path):
    source = tmp_path / "notes.txt"
    source.write_text("Hello world.", encoding="utf-8")
    outputs = tmp_path / "outputs"

    result = create_branded_pdf(source, outputs)
    text = PdfReader(result.output_path).pages[0].extract_text()

    assert "Almond Financial | Internal Use" in text
    assert "Page 1 of" in text


def test_fast_mode_uses_streaming_page_numbers_and_reports_progress(tmp_path):
    source = tmp_path / "notes.txt"
    source.write_text("Hello world.", encoding="utf-8")
    progress = []

    result = create_branded_pdf(source, tmp_path / "outputs", fast=True, progress=progress.append)
    text = PdfReader(result.output_path).pages[0].extract_text()

    assert "Page 1" in text
    assert "Page 1 of" not in text
    assert "Reading document…" in progress
    assert "Rendering streaming PDF…" in progress
    assert "Verifying output…" in progress


def test_pdf_page_selection_limits_source_content(tmp_path):
    from reportlab.pdfgen import canvas as rl_canvas

    source = tmp_path / "source.pdf"
    canvas = rl_canvas.Canvas(str(source))
    for page_number in range(1, 4):
        canvas.drawString(72, 700, f"Unique source page {page_number}")
        canvas.showPage()
    canvas.save()

    result = create_branded_pdf(
        source,
        tmp_path / "outputs",
        fast=True,
        page_range=(1, 2),
    )
    text = " ".join(page.extract_text() for page in PdfReader(result.output_path).pages)

    assert "source page 1" in text
    assert "source page 2" in text
    assert "source page 3" not in text
    assert result.warnings == ["Included source pages 1–2 of 3."]
    assert "EXCERPT" in text
    assert "Included source pages 1" in text


def test_collision_safe_filenames(tmp_path):
    source = tmp_path / "notes.txt"
    source.write_text("Hello world.", encoding="utf-8")
    outputs = tmp_path / "outputs"

    first = create_branded_pdf(source, outputs)
    second = create_branded_pdf(source, outputs)
    third = create_branded_pdf(source, outputs)

    assert first.output_path.endswith("notes-almond.pdf")
    assert second.output_path.endswith("notes-almond-2.pdf")
    assert third.output_path.endswith("notes-almond-3.pdf")
    assert len(set([first.output_path, second.output_path, third.output_path])) == 3


def test_title_override(tmp_path):
    source = tmp_path / "notes.txt"
    source.write_text("Body text.", encoding="utf-8")
    outputs = tmp_path / "outputs"

    result = create_branded_pdf(source, outputs, title_override="Custom Title")
    text = PdfReader(result.output_path).pages[0].extract_text()

    assert "Custom Title" in text


def test_output_filename_sanitizer_neutralises_traversal_attempts():
    # The sanitizer strips path separators from a source stem before it is
    # ever combined with the outputs directory, so a crafted filename can't
    # be used to escape outputs/.
    from app.pdf.service import _sanitize_stem

    assert "/" not in _sanitize_stem("../../etc/passwd")
    assert "\\" not in _sanitize_stem("..\\..\\windows\\system32")


def test_unsupported_source_raises_document_error(tmp_path):
    source = tmp_path / "sheet.xlsx"
    source.write_bytes(b"not real")
    outputs = tmp_path / "outputs"

    with pytest.raises(DocumentError, match="Unsupported document type"):
        create_branded_pdf(source, outputs)

    assert not outputs.exists() or not list(outputs.iterdir())


def test_missing_source_raises_document_error(tmp_path):
    outputs = tmp_path / "outputs"

    with pytest.raises(DocumentError, match="not found"):
        create_branded_pdf(tmp_path / "missing.txt", outputs)


def test_table_block_renders_without_error(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_heading("Report", level=1)
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Revenue"
    table.rows[1].cells[1].text = "100"
    source = tmp_path / "report.docx"
    doc.save(source)
    outputs = tmp_path / "outputs"

    result = create_branded_pdf(source, outputs)
    text = PdfReader(result.output_path).pages[0].extract_text()

    assert "Metric" in text
    assert "Revenue" in text


def test_failed_render_leaves_no_partial_output(tmp_path, monkeypatch):
    source = tmp_path / "notes.txt"
    source.write_text("Hello world.", encoding="utf-8")
    outputs = tmp_path / "outputs"

    def boom(*_args, **_kwargs):
        raise PdfGenerationError("simulated render failure")

    monkeypatch.setattr("app.pdf.service.render_pdf", boom)

    with pytest.raises(PdfGenerationError, match="simulated render failure"):
        create_branded_pdf(source, outputs)

    assert not (outputs / "notes-almond.pdf").exists()
